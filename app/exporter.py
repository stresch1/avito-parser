from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .columns import COLUMNS
from .paths import data_dir

EXPORTS_DIR = data_dir() / "exports"


def export_rows(rows: list[dict], filename: str, columns: list[str] | None = None) -> Path:
    EXPORTS_DIR.mkdir(parents=True, exist_ok=True)
    cols = columns or COLUMNS

    wb = Workbook()
    ws = wb.active
    ws.title = "Объявления"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2E7D32", end_color="2E7D32", fill_type="solid")

    for col_idx, col_name in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, col_name in enumerate(cols, start=1):
            ws.cell(row=row_idx, column=col_idx, value=row.get(col_name, ""))

    for col_idx, col_name in enumerate(cols, start=1):
        letter = get_column_letter(col_idx)
        width = min(max(len(col_name) + 4, 12), 60)
        ws.column_dimensions[letter].width = width

    ws.freeze_panes = "A2"

    out_path = EXPORTS_DIR / filename
    wb.save(out_path)
    return out_path
